import sys
import json
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Expects input_md, output_csv, output_xlsx as arguments
if len(sys.argv) != 4:
    print("Usage: python3 report_parser.py <input_md> <output_csv> <output_xlsx>")
    sys.exit(1)

MD_FILE, CSV_FILE, XLSX_FILE = sys.argv[1], sys.argv[2], sys.argv[3]

with open(MD_FILE, "r", encoding="utf-8") as f:
    content = f.read()

# Strip ANSI colors just in case
content = re.sub(r'\x1B\[[0-?]*[ -/]*[@-~]', '', content)
sections = re.split(r'={10,}', content)
rows = []

for section in sections:
    header_match = re.search(r'INSTANCE\s+#(\d+):\s+(i-[\w\d]+) \((.*?)\)', section)
    if not header_match:
        continue

    num, iid, name = header_match.groups()
    
    # Defaults
    data = {
        "Uptime": "N/A", "Falcon Service": "N/A", "Pending Updates": "N/A",
        "Hostname": "N/A", "Kernel": "N/A", "OS": "N/A"
    }
    os_type = "Unknown"
    
    if "Skipped" in section:
        status = "Skipped"
    elif "Dry-run" in section:
        status = "Dry-Run"
    elif "ERROR" in section:
        status = "Error"
    else:
        status = "Success"
        # Find and parse the JSON block
        json_match = re.search(r'```json\s*([\s\S]+?)\s*```', section)
        if json_match:
            try:
                ssm_data = json.loads(json_match.group(1))
                if 'hostname' in ssm_data:
                    os_type = "Linux"
                    data["Hostname"] = ssm_data.get("hostname", "N/A")
                    data["Kernel"] = ssm_data.get("kernel", "N/A")
                    data["OS"] = ssm_data.get("os", "N/A")
                elif 'Uptime' in ssm_data:
                    os_type = "Windows"
                    data["Uptime"] = ssm_data.get("Uptime", "N/A")
                    data["Falcon Service"] = ssm_data.get("FalconService", "N/A")
                    # Join list of updates into a string
                    updates = ssm_data.get("PendingUpdates", ["N/A"])
                    data["Pending Updates"] = ", ".join(updates)
            except json.JSONDecodeError:
                status = "Error: Invalid JSON"

    rows.append([
        num, iid, name, os_type, status,
        data["Uptime"], data["Falcon Service"], data["Pending Updates"],
        data["Hostname"], data["Kernel"], data["OS"]
    ])

headers = [
    "#", "Instance ID", "Instance Name", "OS Type", "Status",
    "Uptime", "Falcon Service", "Pending Updates", "Hostname", "Kernel", "OS"
]

# Create DataFrame and save to CSV/Excel
df = pd.DataFrame(rows, columns=headers)
df.sort_values(by=["Status", "OS Type", "Instance ID"], inplace=True)
df.to_csv(CSV_FILE, index=False)
df.to_excel(XLSX_FILE, sheet_name="SSM_Report", index=False)

# --- Add Excel table styling and auto-fit columns ---
wb = load_workbook(XLSX_FILE)
ws = wb.active
table_range = f"A1:{chr(ord('A') + len(headers) - 1)}{len(df) + 1}"
table = Table(displayName="ReportData", ref=table_range)
style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
table.tableStyleInfo = style
ws.add_table(table)

for column in ws.columns:
    max_length = 0
    column_letter = column[0].column_letter
    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(cell.value)
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column_letter].width = adjusted_width

wb.save(XLSX_FILE)
print(f"Successfully created CSV: {CSV_FILE}")
print(f"Successfully created Excel: {XLSX_FILE}")